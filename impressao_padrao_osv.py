from reportlab.lib.pagesizes import A4
from reportlab.platypus import Image, SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from datetime import datetime
from conexao import conectar
import os
import re
from tkinter import messagebox

conexao = conectar()

def obter_dados_osv(osv_id, ano):
    if conexao:
        try:
            cursor = conexao.cursor()

            cursor.execute("""
                SELECT id, ano, data_cadastro, modelo, descricao, observacao 
                FROM ordem_servico 
                WHERE id = %s AND ano = %s
            """, (osv_id, ano))
            dados_osv = cursor.fetchone()

            if not dados_osv:
                print(f"Ordem de serviço com ID {osv_id} não encontrada.")
                return None, None, None, None

            cursor.execute("""
                SELECT osp.produto_id, p.nome, osp.sequencial, osp.quantidade, osp.valor_unitario, osp.valor_total 
                FROM ordem_servico_produtos osp
                JOIN produtos p ON p.id = osp.produto_id
                WHERE osp.osv_id = %s AND osp.ano = %s
            """, (osv_id, ano))
            produtos = cursor.fetchall()

            cursor.execute("""
                SELECT oss.servico_id, p.nome, oss.sequencial, oss.quantidade, oss.valor_unitario, oss.valor_total 
                FROM ordem_servico_servicos oss
                JOIN produtos p ON p.id = oss.servico_id
                WHERE oss.osv_id = %s AND oss.ano = %s
            """, (osv_id, ano))
            servicos = cursor.fetchall()

            cursor.execute("""
                SELECT c.nome, c.cpf, ci.nome AS cidade, c.telefone
                FROM clientes c
                JOIN cidades ci ON c.cidade_id = ci.id
                JOIN ordem_servico os ON os.cliente_id = c.id
                WHERE os.id = %s AND os.ano = %s
            """, (osv_id, ano))
            clientes = cursor.fetchall()

            gerar_pdf_osv(dados_osv, produtos, servicos, clientes)
            return dados_osv, produtos, servicos, clientes

        except Exception as e:
            print(f"Erro ao consultar o banco de dados: {e}")
            return None, None, None, None

def formatar_telefone(telefone):
    telefone = re.sub(r'\D', '', telefone)  # Remove qualquer caractere não numérico
    if len(telefone) == 11:  # (XX) XXXXX-XXXX
        return f"({telefone[:2]}) {telefone[2:7]}-{telefone[7:]}"
    elif len(telefone) == 10:  # (XX) XXXX-XXXX
        return f"({telefone[:2]}) {telefone[2:6]}-{telefone[6:]}"
    else:
        return telefone  # Retorna o telefone sem formatação se não for válido
    
# Função para formatar CPF
def formatar_cpf(cpf):
    cpf = re.sub(r'\D', '', cpf)  # Remove qualquer caractere não numérico
    if len(cpf) == 11:  # CPF tem 11 dígitos
        return f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}"
    else:
        return cpf  # Retorna o CPF sem formatação se não for válido    

def gerar_pdf_osv(dados_osv, produtos, servicos, clientes, caminho_arquivo="ordem_servico.pdf"):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    logo_path = os.path.join(script_dir, "logo.png")

    doc = SimpleDocTemplate(caminho_arquivo, pagesize=A4)
    elementos = []

    styles = getSampleStyleSheet()
    estilo_normal = ParagraphStyle(name='Normal', fontSize=12)

    if clientes:
        nome_cliente, cpf, cidade, telefone = clientes[0]
    else:
        return

    numero = dados_osv[0]
    ano = dados_osv[1]
    data = dados_osv[2].strftime("%d/%m/%Y")
    modelo = dados_osv[3]
    descricao = dados_osv[4]
    observacao = dados_osv[5]

    cpf_formatado = formatar_cpf(cpf)
    telefone_formatado = formatar_telefone(telefone)

    # Preenche o arquivo Word com base nos dados da OSV
    preencher_osv_word("OSV.docx", "OSV_Preenchido.docx", {
        "@NUM": numero,
        "@ANO": ano,
        "@NOME_CLIENTE": nome_cliente,
        "@CPF_CLIENTE": cpf_formatado,
        "@CIDADE": cidade,
        "@TELEFONE": telefone_formatado,
        "@MODELO": modelo,
        "@DESCRICAO": descricao,
        "@DATA": data
    })

    try:
        imagem_logo = Image(logo_path, width=400, height=100)
        elementos.append(imagem_logo)
    except Exception as e:
        print(f"Erro ao carregar a logo: {e}")

    elementos.append(Spacer(1, 12))
    elementos.append(Paragraph(f"Ordem de Serviço {numero}", styles['Title']))
    elementos.append(Spacer(1, 12))
    elementos.append(Paragraph(f"Cliente: {nome_cliente}", estilo_normal))
    elementos.append(Paragraph(f"CPF: {cpf}", estilo_normal))
    elementos.append(Paragraph(f"Cidade: {cidade}", estilo_normal))
    elementos.append(Paragraph(f"Telefone: {telefone}", estilo_normal))
    elementos.append(Spacer(1, 12))
    elementos.append(Paragraph(f"Número: {numero}/{ano}", estilo_normal))
    elementos.append(Paragraph(f"Data: {data}", estilo_normal))
    elementos.append(Paragraph(f"Modelo: {modelo}", estilo_normal))
    elementos.append(Paragraph(f"Descrição: {descricao}", estilo_normal))
    elementos.append(Paragraph(f"Observação: {observacao}", estilo_normal))
    elementos.append(Spacer(1, 12))

    elementos.append(Paragraph("Produtos:", estilo_normal))
    for produto in produtos:
        elementos.append(Paragraph(f"Produto: {produto[1]} | Quantidade: {produto[3]} | Unit: R$ {produto[4]:.2f} | Total: R$ {produto[5]:.2f}", estilo_normal))
        elementos.append(Spacer(1, 6))

    elementos.append(Spacer(1, 12))
    elementos.append(Paragraph("Serviços:", estilo_normal))
    for servico in servicos:
        elementos.append(Paragraph(f"Serviço: {servico[1]} | Quantidade: {servico[3]} | Unit: R$ {servico[4]:.2f} | Total: R$ {servico[5]:.2f}", estilo_normal))
        elementos.append(Spacer(1, 6))

    elementos.append(Spacer(1, 12))
    elementos.append(Paragraph(f"Data: {data}", estilo_normal))
    elementos.append(Spacer(1, 12))
    elementos.append(Paragraph("Assinatura do Cliente: ______________________", estilo_normal))
    elementos.append(Spacer(1, 12))
    elementos.append(Paragraph("AV. XV DE NOVEMBRO, 535, SALA 117, CENTRO, JOAÇABA", estilo_normal))
    elementos.append(Paragraph("FONE: (49) 3521-1454 OU (49) 9 9111-2696", estilo_normal))

    doc.build(elementos)
    print(f"PDF gerado com sucesso: {caminho_arquivo}")

def preencher_osv_word(caminho_entrada, caminho_saida, dados):
    """
    Substitui os placeholders em um arquivo .docx e salva o resultado.
    """
    doc = Document(caminho_entrada)

    for paragrafo in doc.paragraphs:
        for chave, valor in dados.items():
            if chave in paragrafo.text:
                for run in paragrafo.runs:
                    if chave in run.text:
                        run.text = run.text.replace(chave, str(valor))

    try:
        doc.save(caminho_saida)
        print(f"Documento preenchido salvo em: {caminho_saida}")
        messagebox.showinfo("Sucesso", "Arquivo gerado com sucesso! Caminho: {}".format(caminho_saida))
    except Exception as e:
        print(f"Ocorreu um erro: {e}")
        messagebox.showerror("Erro", f"Ocorreu um erro ao salvar o arquivo: {e}")


